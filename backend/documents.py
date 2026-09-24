"""Executive ATS-friendly document generation for CV and cover letter exports.
Outputs matching DOCX and PDF formats with clean typography, elegant dividers, and structured layouts.
"""
import os
import re
import shutil
import subprocess
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.graphics.shapes import Drawing, Line

from . import store

PAGE_WIDTH_PT = 595.27
PAGE_HEIGHT_PT = 841.89
LEFT_MARGIN_PT = 38
RIGHT_MARGIN_PT = 38
PRINTABLE_WIDTH_PT = PAGE_WIDTH_PT - LEFT_MARGIN_PT - RIGHT_MARGIN_PT

def _create_divider(width=PRINTABLE_WIDTH_PT, color_hex='#246952', thickness=1.2):
    d = Drawing(width, 4)
    d.add(Line(0, 2, width, 2, strokeColor=colors.HexColor(color_hex), strokeWidth=thickness))
    return d

def generate(profile, package_id):
    folder = store.DATA / 'documents' / package_id
    folder.mkdir(parents=True, exist_ok=True)

    is_cover_letter = any(s.get('title', '').lower() == 'cover letter' for s in profile.get('sections', []))

    # --- 1. DOCX GENERATION ---
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    sec.left_margin = sec.right_margin = Inches(0.55)
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(3.5)
    normal.paragraph_format.line_spacing = 1.1

    doc.styles['Title'].font.name = 'Calibri'
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Title'].font.bold = True
    doc.styles['Title'].font.color.rgb = RGBColor(23, 62, 50)  # Dark emerald

    doc.styles['Heading 1'].font.name = 'Calibri'
    doc.styles['Heading 1'].font.size = Pt(11)
    doc.styles['Heading 1'].font.bold = True
    doc.styles['Heading 1'].font.color.rgb = RGBColor(36, 105, 82)  # Primary accent
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(3)

    p_title = doc.add_paragraph(profile['name'], 'Title')
    p_title.paragraph_format.space_after = Pt(2)

    p_head = doc.add_paragraph(profile['headline'])
    p_head.paragraph_format.space_after = Pt(4)
    p_head.runs[0].font.size = Pt(10.5)
    p_head.runs[0].font.color.rgb = RGBColor(74, 85, 104)

    contact = '  |  '.join(x for x in [profile.get('location'), profile.get('phone'), profile.get('email')] if x)
    if contact:
        p_contact = doc.add_paragraph(contact)
        p_contact.paragraph_format.space_after = Pt(2)
        p_contact.runs[0].font.size = Pt(9)
        p_contact.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    if profile.get('links'):
        p_links = doc.add_paragraph('  |  '.join(profile['links']))
        p_links.paragraph_format.space_after = Pt(8)
        p_links.runs[0].font.size = Pt(9)
        p_links.runs[0].font.color.rgb = RGBColor(43, 108, 176)

    for section in profile['sections']:
        doc.add_paragraph(section['title'].upper(), 'Heading 1')
        for index, item in enumerate(section['items']):
            text = item['text'].strip()
            p = doc.add_paragraph()
            if is_cover_letter:
                p.add_run(text)
                p.paragraph_format.space_after = Pt(6)
            elif re.search(r'\d{2}/\d{4}|Present\b|Freelance\b|Internship\b', text):
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(4)
                run = p.add_run(text)
                run.bold = True
                run.font.color.rgb = RGBColor(30, 41, 59)
            elif section['title'] == 'Technical Skills' and ':' in text:
                cat, val = text.split(':', 1)
                r1 = p.add_run(cat + ':')
                r1.bold = True
                p.add_run(val)
                p.paragraph_format.space_after = Pt(2)
            else:
                # Standard bullet point
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.first_line_indent = Inches(-0.12)
                p.add_run('• ' + text)
                p.paragraph_format.space_after = Pt(2.5)

    doc.save(folder / 'cv.docx')

    # --- 2. PDF GENERATION (ReportLab & Optional LibreOffice) ---
    office = os.getenv('SOFFICE_PATH') or shutil.which('soffice')
    method = 'reportlab'
    if office:
        try:
            subprocess.run([
                office,
                '-env:UserInstallation=' + (folder / 'lo-profile').resolve().as_uri(),
                '--headless',
                '--convert-to', 'pdf',
                '--outdir', str(folder),
                str(folder / 'cv.docx')
            ], check=True, timeout=60, capture_output=True)
            if (folder / 'cv.pdf').exists():
                method = 'libreoffice'
        except (OSError, subprocess.SubprocessError):
            pass

    if method == 'reportlab':
        name_style = ParagraphStyle(
            'CandidateName',
            fontName='Helvetica-Bold',
            fontSize=21,
            leading=25,
            textColor=colors.HexColor('#173e32'),
            spaceAfter=3
        )
        headline_style = ParagraphStyle(
            'Headline',
            fontName='Helvetica-Bold',
            fontSize=10.5,
            leading=14,
            textColor=colors.HexColor('#2d4f40'),
            spaceAfter=4
        )
        meta_style = ParagraphStyle(
            'ContactMeta',
            fontName='Helvetica',
            fontSize=8.8,
            leading=12,
            textColor=colors.HexColor('#4a5568'),
            spaceAfter=2
        )
        links_style = ParagraphStyle(
            'LinksMeta',
            parent=meta_style,
            textColor=colors.HexColor('#246952'),
            spaceAfter=6
        )
        heading_style = ParagraphStyle(
            'SectionHeader',
            fontName='Helvetica-Bold',
            fontSize=10.5,
            leading=13,
            textColor=colors.HexColor('#173e32'),
            spaceBefore=8,
            spaceAfter=2,
            keepWithNext=True
        )
        body_style = ParagraphStyle(
            'StandardBody',
            fontName='Helvetica',
            fontSize=9.2,
            leading=13,
            textColor=colors.HexColor('#1f2937'),
            spaceAfter=3.5
        )
        bullet_style = ParagraphStyle(
            'BulletPoint',
            parent=body_style,
            leftIndent=12,
            firstLineIndent=-7,
            spaceAfter=3
        )
        role_header_style = ParagraphStyle(
            'RoleHeader',
            fontName='Helvetica-Bold',
            fontSize=9.5,
            leading=13.5,
            textColor=colors.HexColor('#0f172a'),
            spaceBefore=5,
            spaceAfter=2,
            keepWithNext=True
        )
        cover_body_style = ParagraphStyle(
            'CoverLetterBody',
            fontName='Helvetica',
            fontSize=10,
            leading=15,
            textColor=colors.HexColor('#1f2937'),
            spaceAfter=8,
            alignment=TA_LEFT
        )

        story = [
            Paragraph(escape(profile['name']), name_style),
            Paragraph(escape(profile['headline']), headline_style)
        ]

        if contact:
            story.append(Paragraph(escape(contact), meta_style))
        if profile.get('links'):
            story.append(Paragraph(escape('   |   '.join(profile['links'])), links_style))

        # Divider under top contact ribbon
        story.append(_create_divider(width=PRINTABLE_WIDTH_PT, color_hex='#c2dbca', thickness=0.8))
        story.append(Spacer(1, 4))

        for section in profile['sections']:
            sec_title = section['title'].strip()
            story.append(Paragraph(escape(sec_title.upper()), heading_style))
            story.append(_create_divider(width=PRINTABLE_WIDTH_PT, color_hex='#246952', thickness=1.0))
            story.append(Spacer(1, 3))

            group = []
            for item in section['items']:
                raw = item['text'].strip()
                if not raw:
                    continue

                if is_cover_letter:
                    story.append(Paragraph(escape(raw), cover_body_style))
                elif sec_title == 'Summary':
                    story.append(Paragraph(escape(raw), body_style))
                elif sec_title == 'Technical Skills' and ':' in raw:
                    cat, vals = raw.split(':', 1)
                    styled_skill = f"<b>{escape(cat)}:</b>{escape(vals)}"
                    story.append(Paragraph(styled_skill, body_style))
                elif re.search(r'\d{2}/\d{4}|Present\b|Freelance\b|Internship\b', raw):
                    if group:
                        story.append(KeepTogether(group))
                        group = []
                    # Emphasize role header
                    styled_role = f"<b>{escape(raw)}</b>"
                    role_p = Paragraph(styled_role, role_header_style)
                    if sec_title == 'Professional Experience':
                        group.append(role_p)
                    else:
                        story.append(role_p)
                else:
                    # Formatted bullet point
                    bullet_text = f"&bull;&nbsp;&nbsp;{escape(raw)}"
                    bp = Paragraph(bullet_text, bullet_style)
                    if sec_title == 'Professional Experience':
                        group.append(bp)
                    else:
                        story.append(bp)

            if group:
                story.append(KeepTogether(group))

            story.append(Spacer(1, 4))

        doc_pdf = SimpleDocTemplate(
            str(folder / 'cv.pdf'),
            pagesize=A4,
            leftMargin=LEFT_MARGIN_PT,
            rightMargin=RIGHT_MARGIN_PT,
            topMargin=34,
            bottomMargin=34,
            title=f"{profile['name']} - CV",
            author=profile['name'],
            subject='Curriculum Vitae'
        )
        doc_pdf.build(story)

    import hashlib
    return {
        fmt: {
            'path': str((folder / f'cv.{fmt}').relative_to(store.DATA)),
            'sha256': hashlib.sha256((folder / f'cv.{fmt}').read_bytes()).hexdigest()
        } for fmt in ('docx', 'pdf')
    } | {'pdf_method': method}

def verify_files(package):
    import hashlib
    for fmt in ('pdf', 'docx', *[k for k in ('cover_letter_pdf', 'cover_letter_docx') if k in package['files']]):
        info = package['files'].get(fmt)
        if not info:
            raise ValueError('CV export is missing. Prepare the package again.')
        path = (store.DATA / info['path']).resolve()
        if not path.is_relative_to(store.DATA.resolve()) or not path.exists() or hashlib.sha256(path.read_bytes()).hexdigest() != info['sha256']:
            raise ValueError('CV file changed after preparation. Prepare and approve a new package.')
